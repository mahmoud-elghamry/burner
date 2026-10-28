#!/usr/bin/env python3
"""Build an Openness SimaticML PLC tag table from the generated C&E workbook."""

from __future__ import annotations

import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "deliverables" / "PH-LCP-01-CE_RevA.xlsx"
OUTPUT = ROOT / "build" / "tags" / "IO.xml"
NAME_RE = re.compile(r"^[A-Za-z][A-Za-z0-9]*(_[A-Za-z0-9]+)*$")
BIT_ADDRESS_RE = re.compile(r"^([IQ])(\d+)\.(\d+)$")
WORD_ADDRESS_RE = re.compile(r"^IW(\d+)$")
REQUIRED_HEADERS = ("Address", "Tag", "Description", "PLC Tag")


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def validate_name(name: str, row_number: int) -> None:
    if not NAME_RE.fullmatch(name):
        raise ValueError(
            f"row {row_number}: PLC tag {name!r} breaks the 0005 naming pattern"
        )
    function_parts = name.split("_")[1:]
    if any(part.islower() for part in function_parts):
        raise ValueError(
            f"row {row_number}: PLC tag {name!r} has a lowercase-only Function part"
        )


def parse_address(address: str, row_number: int) -> tuple[str, str]:
    match = BIT_ADDRESS_RE.fullmatch(address)
    if match:
        area, byte_text, bit_text = match.groups()
        byte, bit = int(byte_text), int(bit_text)
        if bit > 7:
            raise ValueError(f"row {row_number}: invalid bit address {address!r}")
        return f"%{area}{byte}.{bit}", "Bool"

    match = WORD_ADDRESS_RE.fullmatch(address)
    if match:
        return f"%IW{int(match.group(1))}", "Int"

    raise ValueError(f"row {row_number}: cannot parse I/O address {address!r}")


def find_header(ws) -> tuple[int, dict[str, int]]:
    for row_number, values in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True), 1
    ):
        columns = {
            text(value): column_number
            for column_number, value in enumerate(values, 1)
            if text(value)
        }
        if all(header in columns for header in REQUIRED_HEADERS):
            return row_number, columns
    raise ValueError(
        f"workbook sheet {ws.title!r} is missing headers: "
        + ", ".join(REQUIRED_HEADERS)
    )


def read_tags() -> list[dict[str, str]]:
    if not WORKBOOK.is_file():
        raise FileNotFoundError(f"source workbook not found: {WORKBOOK}")

    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        if "IO List" not in wb.sheetnames:
            raise ValueError("source workbook has no 'IO List' sheet")
        ws = wb["IO List"]
        header_row, columns = find_header(ws)
        tags: list[dict[str, str]] = []
        names: dict[str, int] = {}
        addresses: dict[str, tuple[str, int]] = {}

        for row_number, values in enumerate(
            ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
        ):
            def cell(header: str) -> str:
                index = columns[header] - 1
                return text(values[index]) if index < len(values) else ""

            name = cell("PLC Tag")
            if not name:
                continue

            address = cell("Address")
            drawing_tag = cell("Tag")
            comment = cell("Description")
            if not address:
                raise ValueError(f"row {row_number}: PLC tag {name!r} has no address")
            if not drawing_tag:
                raise ValueError(f"row {row_number}: PLC tag {name!r} has no drawing tag")
            validate_name(name, row_number)

            if name in names:
                raise ValueError(
                    f"row {row_number}: duplicate PLC tag name {name!r} "
                    f"(first used on row {names[name]})"
                )
            names[name] = row_number

            logical_address, data_type = parse_address(address, row_number)
            if logical_address in addresses:
                old_name, old_row = addresses[logical_address]
                raise ValueError(
                    f"row {row_number}: duplicate address {logical_address} for {name!r}; "
                    f"already used by {old_name!r} on row {old_row}"
                )
            addresses[logical_address] = (name, row_number)
            tags.append(
                {
                    "address": logical_address,
                    "drawing_tag": drawing_tag,
                    "name": name,
                    "type": data_type,
                    "comment": comment,
                }
            )

        return tags
    finally:
        wb.close()


def build_xml(tags: list[dict[str, str]]) -> ET.ElementTree:
    document = ET.Element("Document")
    ET.SubElement(document, "Engineering", {"version": "V15.1"})

    info = ET.SubElement(document, "DocumentInfo")
    created = datetime.now(timezone.utc).isoformat(timespec="microseconds").replace(
        "+00:00", "Z"
    )
    ET.SubElement(info, "Created").text = created
    ET.SubElement(info, "ExportSetting").text = "WithDefaults"
    products = ET.SubElement(info, "InstalledProducts")
    for kind, display_name in (
        ("Product", "Totally Integrated Automation Portal"),
        ("OptionPackage", "TIA Portal Openness"),
        ("Product", "STEP 7 Professional"),
        ("Product", "WinCC Professional"),
    ):
        product = ET.SubElement(products, kind)
        ET.SubElement(product, "DisplayName").text = display_name
        ET.SubElement(product, "DisplayVersion").text = "V15.1"

    table = ET.SubElement(document, "SW.Tags.PlcTagTable", {"ID": "0"})
    attributes = ET.SubElement(table, "AttributeList")
    ET.SubElement(attributes, "Name").text = "IO"
    object_list = ET.SubElement(table, "ObjectList")

    next_id = 1
    for item in tags:
        tag = ET.SubElement(
            object_list,
            "SW.Tags.PlcTag",
            {"ID": str(next_id), "CompositionName": "Tags"},
        )
        next_id += 1
        tag_attributes = ET.SubElement(tag, "AttributeList")
        ET.SubElement(tag_attributes, "DataTypeName").text = item["type"]
        ET.SubElement(tag_attributes, "LogicalAddress").text = item["address"]
        ET.SubElement(tag_attributes, "Name").text = item["name"]

        tag_objects = ET.SubElement(tag, "ObjectList")
        comment = ET.SubElement(
            tag_objects,
            "MultilingualText",
            {"ID": str(next_id), "CompositionName": "Comment"},
        )
        next_id += 1
        comment_objects = ET.SubElement(comment, "ObjectList")
        item_text = ET.SubElement(
            comment_objects,
            "MultilingualTextItem",
            {"ID": str(next_id), "CompositionName": "Items"},
        )
        next_id += 1
        item_attributes = ET.SubElement(item_text, "AttributeList")
        ET.SubElement(item_attributes, "Culture").text = "en-US"
        ET.SubElement(item_attributes, "Text").text = item["comment"]

    return ET.ElementTree(document)


def main() -> int:
    tags = read_tags()
    tree = build_xml(tags)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    ET.indent(tree, space="  ")
    tree.write(OUTPUT, encoding="utf-8", xml_declaration=True)
    print(f"wrote {OUTPUT.relative_to(ROOT)} ({len(tags)} tags)")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
