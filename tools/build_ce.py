import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "deliverables", "PH-LCP-01-CE_RevA.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

F = "Arial"
thin = Side(style="thin", color="808080")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
GRP = PatternFill("solid", fgColor="EDEDED")
INPUT = PatternFill("solid", fgColor="FFFF99")
FILL = {
    "STOP": PatternFill("solid", fgColor="F8CBAD"),
    "RUN": PatternFill("solid", fgColor="C6EFCE"),
    "OPEN": PatternFill("solid", fgColor="FFE699"),
    "CLOSE": PatternFill("solid", fgColor="FFE699"),
    "ALARM": PatternFill("solid", fgColor="FFF2CC"),
    "LAMP": PatternFill("solid", fgColor="DDEBF7"),
}
SRC_FILL = {
    "Client": PatternFill("solid", fgColor="E2EFDA"),
    "PH proposal": PatternFill("solid", fgColor="FCE4D6"),
}


def font(**kw):
    kw.setdefault("name", F)
    kw.setdefault("size", 9)
    return Font(**kw)


def put(ws, r, c, v, bold=False, fill=None, color="000000", wrap=True, h="left", size=9, border=True):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font(bold=bold, color=color, size=size)
    cell.alignment = Alignment(wrap_text=wrap, vertical="center", horizontal=h)
    if fill:
        cell.fill = fill
    if border:
        cell.border = BOX
    return cell


def title_block(ws, title, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    put(ws, 1, 1, title, bold=True, fill=HDR, color="FFFFFF", size=12, border=False)
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    put(ws, 2, 1, "Project: Burner PetroBaker  |  Panel: LCP-01  |  Ref. drawing: PH-LCP-01 V1  |  "
                  "Document: PH-LCP-01-CE  |  Rev. A (for client review)  |  26-10-2026",
        size=9, border=False)


wb = Workbook()

# ------------------------------------------------------------------ Cover
ws = wb.active
ws.title = "Cover"
title_block(ws, "CAUSE & EFFECT MATRIX — LCP-01 Scrubber & Burner Control", 4)
rows = [
    ("Purpose", "Defines the PLC (LCP-01) response to every process and safety cause, as requested in the "
                "client e-mail of 04-Oct-2026 (Eng. Amr Farag)."),
    ("Scope", "Gas scrubber alarms/trips, ESD, burner BR01 (Riello RS100) and BR02 (DLX 53 SHFP) "
              "trips and permissives, tank temperature control (client sheet 'Temperature control of burner'). "
              "This matrix lists causes and effects only. Operating modes (AUTO / MAN), START / STOP, "
              "RESET, start-up after power loss and lamps are described in the Control Narrative "
              "(docs/التشغيل.md)."),
    ("Not in scope (Rev. A)", "Communication with the Master PLC in the control room — deferred to a later "
                              "revision.\nScrubber low level switch — not installed (DI I12.5 reserved)."),
    ("Burner flame safety", "Purge, ignition, flame supervision and the pilot / main gas valves are handled by "
                            "each burner's own certified control box (RS100: RMG88; DLX 53: SQ DUAL). "
                            "LCP-01 only gives the burner its heat request (TL / L) and stage request "
                            "(TR / T2), and reads its RUN and LOCKOUT signals. LCP-01 never drives a gas "
                            "valve directly."),
    ("Hardwired safety", "E-STOP S0 opens the 230 V control supply of BR01 (NC2) and BR02 (NC3) directly, "
                         "independent of the PLC. Safety thermostat TS (manual reset) is in series with the "
                         "BR01 control supply. These act even if the PLC fails."),
    ("Fail-safe outputs", "All PLC outputs de-energise on PLC STOP or 24 VDC loss. Valve types were chosen so "
                          "that de-energised = safe: XV-401 flare (NO) OPENS, XV-402 drain (NO) OPENS, "
                          "XV-403 suction (NC) CLOSES, both burners lose their heat request."),
    ("How to read the matrix", "One row per cause. Cells under each effect say what happens to that output "
                               "while the cause is active. Empty = no action.\n"
                               "STOP = output de-energised / burner request removed.  RUN = energised.\n"
                               "OPEN / CLOSE = resulting valve position.  ALARM = HMI alarm (the burner fault lamps H3 / H4 are wired to the burner lockout, not to the PLC).\n"
                               "Column 'Source' shows where each line comes from: Client = client document; "
                               "PH proposal = proposed by PH, to be approved by the client."),
    ("Requested from client", "On sheet 'C&E Matrix', fill the two yellow columns for every row: "
                              "'Client response' (Approved / Approved with comment / Rejected) and "
                              "'Client comment'. Rows marked 'PH proposal' need particular attention. "
                              "All other cells are protected against accidental edits (no password). "
                              "Use the filter on the header row, e.g. Source = PH proposal. "
                              "Parameter values are on sheet 'Parameters' and are adjustable from the HMI."),
]
r = 4
for k, v in rows:
    put(ws, r, 1, k, bold=True, fill=SUB)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    put(ws, r, 2, v)
    ws.row_dimensions[r].height = 13 * (len(v) // 125 + 1 + v.count("\n")) + 8
    r += 1
r += 1
put(ws, r, 1, "Revision", bold=True, fill=SUB); put(ws, r, 2, "Date", bold=True, fill=SUB)
put(ws, r, 3, "Description", bold=True, fill=SUB); put(ws, r, 4, "By / Approved", bold=True, fill=SUB)
put(ws, r + 1, 1, "A"); put(ws, r + 1, 2, "26-10-2026")
put(ws, r + 1, 3, "First issue for client review"); put(ws, r + 1, 4, "PH / Eng. Ahmed Zain")
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 60
ws.column_dimensions["D"].width = 40

# ------------------------------------------------------------------ Matrix
ws = wb.create_sheet("C&E Matrix")
effects = [
    ("E1", "BR01 heat request (TL)", "Q0.0 / R1"),
    ("E2", "BR01 2nd stage (TR)", "Q0.1 / R2"),
    ("E3", "BR02 start (L)", "Q0.5 / R6"),
    ("E4", "BR02 high fire (T2)", "Q0.6 / R7"),
    ("E5", "XV-401 Flare SV (NO)", "Q0.2 / R3"),
    ("E6", "XV-402 Drain SV (NO)", "Q0.3 / R4"),
    ("E7", "XV-403 Suction SV (NC)", "Q0.4 / R5"),
    ("E8", "HMI alarm", "HMI"),
]
cause_cols = ["No.", "Tag", "Cause", "PLC address", "Condition / setpoint", "Type"]
tail_cols = ["Reset", "Source", "Remarks", "Client response", "Client comment"]
NC = len(cause_cols)
ncols = NC + len(effects) + len(tail_cols)
title_block(ws, "CAUSE & EFFECT MATRIX", ncols)

hr = 4
for i, h in enumerate(cause_cols, 1):
    ws.merge_cells(start_row=hr, start_column=i, end_row=hr + 2, end_column=i)
    put(ws, hr, i, h, bold=True, fill=HDR, color="FFFFFF", h="center")
ws.merge_cells(start_row=hr, start_column=NC + 1, end_row=hr, end_column=NC + len(effects))
put(ws, hr, NC + 1, "EFFECTS", bold=True, fill=HDR, color="FFFFFF", h="center")
for j, (eid, name, addr) in enumerate(effects):
    c = NC + 1 + j
    put(ws, hr + 1, c, f"{eid}\n{name}", bold=True, fill=SUB, h="center")
    put(ws, hr + 2, c, addr, fill=SUB, h="center")
for k, h in enumerate(tail_cols):
    c = NC + len(effects) + 1 + k
    ws.merge_cells(start_row=hr, start_column=c, end_row=hr + 2, end_column=c)
    put(ws, hr, c, h, bold=True, fill=HDR, color="FFFFFF", h="center")
ws.row_dimensions[hr + 1].height = 42

SCRUB_TRIP = {"E5": "OPEN", "E6": "OPEN", "E7": "CLOSE", "E8": "ALARM"}
BR1_STOP = {"E1": "STOP", "E2": "STOP"}
BR2_STOP = {"E3": "STOP", "E4": "STOP"}

CL = "Client"
PH = "PH proposal"


def burner_rows(n, br, tk, lsl, te_in, te_out, rr, flt, stop_fx, stg):
    p = f"C{n:02d}"
    te = (f"{te_in[0]} / {te_out[0]}", f"{te_in[1]} / {te_out[1]}")
    alarm = {**stop_fx, "E8": "ALARM"}
    return [
        ("group", f"BURNER {br} — tank {tk}   (operating modes, START / STOP: see Control Narrative)"),
        (f"{p}a", lsl[0], f"{tk} level below low level switch", lsl[1],
         "LSL not reached (contact open)", "Permissive", stop_fx, "Automatic when level returns", CL,
         "Client sheet: burner may run only when level has reached LSL. Contact closes with liquid "
         "present, so a broken wire also blocks the burner."),
        (f"{p}b", te[0], f"{tk} oil temperature at setpoint", te[1], "max(T) >= SP (70 °C)", "Control",
         stop_fx, "Automatic at max(T) <= SP − HYST (65 °C)", CL,
         "Client cases 2/3/6. The higher of inlet/outlet is used and the 5 °C restart band are PH "
         "proposals: each restart costs ~50 s of purge."),
        (f"{p}c", te[0], f"{tk} oil temperature low — fast heat-up", te[1], "max(T) < SP − STG (60 °C)",
         "Control", {stg: "RUN"}, "Automatic at max(T) >= 60 °C", PH,
         "2nd stage / high fire. Between 60 and 70 °C the burner runs on 1st stage."),
        (f"{p}d", te[0], "Thermocouple fault — either of the two", te[1],
         "Wire break or value out of range", "Trip", alarm, "S2 after fault cleared", PH,
         "Either sensor failing stops the burner: without both readings the higher value cannot be trusted."),
        (f"{p}e", flt[0], f"{br} lockout (flame / ignition / air pressure failure)", flt[1],
         "Lockout signal present", "Trip", alarm, "At burner control box, then S2", PH,
         "The burner's own control box has already closed the gas valves. PLC removes the heat request "
         "so the burner does not restart by itself after a local reset."),
        (f"{p}f", rr[0], f"{br} failed to start", rr[1],
         "Heat request ON, no RUN within T_START (120 s)", "Trip", alarm, "S2", PH,
         "RUN = gas valves energised (firing). Normal start ~55 s (RS100 §6.8.1) / ~65 s (DLX 53)."),
        (f"{p}g", rr[0], f"{br} stopped by itself while running", rr[1],
         "RUN lost, no lockout, heat request still ON", "Trip", alarm, "S2", PH,
         "E.g. safety thermostat TS, gas pressure switch, burner supply lost. Without this row the PLC "
         "would keep requesting heat from a burner that is no longer firing."),
    ]


matrix = [
    ("group", "EMERGENCY SHUTDOWN"),
    ("C01", "S0", "E-STOP pressed", "I1.0", "NC1 contact open", "Trip",
     {**BR1_STOP, **BR2_STOP, **SCRUB_TRIP}, "Release S0, then S2; burners need START", PH,
     "Burners are also cut by hardwired S0 NC2/NC3 (client e-mail: ESD impact on burner ON/OFF). "
     "Scrubber valves to safe position on ESD = PH proposal."),
    ("C02", "—", "PLC STOP, 24 VDC loss or power-up", "—", "All outputs de-energised", "Inherent",
     {**BR1_STOP, **BR2_STOP, "E5": "OPEN", "E6": "OPEN", "E7": "CLOSE"},
     "S2 for scrubber; burners need START", PH,
     "Result of valve types (NO / NO / NC) and relay outputs. After power returns nothing restarts "
     "by itself."),
    ("group", "GAS SCRUBBER"),
    ("C03", "PSH-101", "Scrubber high pressure", "I0.0", "Contact open (fail-safe)", "Trip",
     SCRUB_TRIP, "S2 after cleared", CL,
     "Client e-mail: scrubber alarms/trips. Burners are NOT tripped (decision PH / Eng. Ahmed Zain); "
     "if gas pressure to the burner drops, the burner's own gas pressure switch stops it (C07g / C08g)."),
    ("C04", "LSH-102", "Scrubber high level", "I0.1", "Contact open (fail-safe)", "Alarm",
     {"E8": "ALARM"}, "Automatic when cleared", CL, "Pre-alarm before LSHH-103."),
    ("C05", "LSHH-103", "Scrubber high-high level", "I0.2", "Contact open (fail-safe)", "Trip",
     SCRUB_TRIP, "S2 after cleared", CL, "Burners are NOT tripped (as C03)."),
    ("C06", "LSL (scrubber)", "Scrubber low level", "I12.5 (reserved)", "—", "Not installed", {}, "—",
     CL, "Not installed — excluded from Rev. A. Input reserved for future use."),
]
matrix += burner_rows(7, "BR01", "TK01", ("LSL-TK01", "I1.2"), ("TE-501", "IW96"), ("TE-502", "IW98"),
                      ("RR", "I0.3"), ("Rflt", "I0.5"), BR1_STOP, "E2")
matrix.append(("C07h", "—", "BR01 running continuously for 24 h", "—", "Run timer >= 24 h", "Control",
               BR1_STOP, "Automatic (restarts on next demand)", PH,
               "RS100 is approved for intermittent operation and must stop at least once every 24 h "
               "(RS100 manual §5.10)."))
matrix += burner_rows(8, "BR02", "TK02", ("LSL-TK02", "I1.3"), ("TE-503", "IW100"), ("TE-504", "IW102"),
                      ("RR2", "I1.4"), ("Rflt2", "I1.5"), BR2_STOP, "E4")

resp = DataValidation(type="list", formula1='"Approved,Approved with comment,Rejected"', allow_blank=True)
resp.error = "Choose Approved, Approved with comment or Rejected"
resp.prompt = "Approved / Approved with comment / Rejected"
ws.add_data_validation(resp)
r = hr + 3
eidx = {e[0]: NC + 1 + j for j, e in enumerate(effects)}
for row in matrix:
    if row[0] == "group":
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        put(ws, r, 1, row[1], bold=True, fill=GRP)
        r += 1
        continue
    no, tag, cause, addr, cond, typ, fx, reset, src, rem = row
    for i, v in enumerate([no, tag, cause, addr, cond, typ], 1):
        put(ws, r, i, v, h="center" if i in (1, 4, 6) else "left")
    for eid, c in eidx.items():
        v = fx.get(eid)
        put(ws, r, c, v, bold=bool(v), fill=FILL.get(v) if v else None, h="center")
    base = NC + len(effects)
    put(ws, r, base + 1, reset)
    put(ws, r, base + 2, src, fill=SRC_FILL.get(src), h="center")
    put(ws, r, base + 3, rem)
    # the only cells the client edits: unlocked, highlighted, dropdown on the response
    for c in (base + 4, base + 5):
        cell = put(ws, r, c, None, fill=INPUT, h="center" if c == base + 4 else "left")
        cell.protection = Protection(locked=False)
    resp.add(ws.cell(row=r, column=base + 4))
    ws.row_dimensions[r].height = max(28, 12 * (len(rem) // 55 + 1))
    r += 1

widths = [6, 13, 30, 11, 22, 10] + [11] * len(effects) + [20, 12, 50, 16, 40]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = f"D{hr + 3}"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = f"{hr}:{hr + 2}"
ws.auto_filter.ref = f"A{hr + 2}:{get_column_letter(ncols)}{r - 1}"
# no password: this only stops accidental edits; filtering stays allowed
ws.protection.sheet = True
ws.protection.autoFilter = False
ws.protection.sort = False
ws.protection.formatColumns = False
ws.protection.formatRows = False

# ------------------------------------------------------------------ Parameters
ws = wb.create_sheet("Parameters")
title_block(ws, "ADJUSTABLE PARAMETERS (HMI)", 6)
hdr = ["Parameter", "Description", "Default", "Unit", "Range", "Source"]
for i, h in enumerate(hdr, 1):
    put(ws, 4, i, h, bold=True, fill=HDR, color="FFFFFF", h="center")
params = [
    ("SP", "Tank temperature setpoint — burner stops at T >= SP", 70, "°C", "40 – 90",
     "Client sheet (adjustable from SCADA/HMI)"),
    ("HYST", "Restart hysteresis — burner restarts at T <= SP − HYST", 5, "°C", "3 – 10",
     "PH proposal: > type-K accuracy (±2.2 °C), limits short-cycling"),
    ("STG", "2nd stage / high fire while T < SP − STG", 10, "°C", "5 – 30", "PH proposal"),
    ("T_START", "Max. time from heat request to RUN feedback", 120, "s", "90 – 300",
     "PH proposal: ~2 × burner start cycle"),
    ("T_MAXRUN", "Forced stop of BR01 after continuous run", 24, "h", "fixed",
     "RS100 manual §5.10"),
]
for k, p in enumerate(params, 5):
    for i, v in enumerate(p, 1):
        c = put(ws, k, i, v, h="center" if i in (3, 4, 5) else "left")
        if i == 3:
            c.font = font(color="0000FF", bold=True)
put(ws, 11, 1, "Blue values are defaults; they can be changed from the HMI by an authorised user.",
    border=False)
ws.merge_cells("A11:F11")
for i, w in enumerate([12, 55, 10, 8, 12, 50], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ------------------------------------------------------------------ I/O list
ws = wb.create_sheet("IO List")
title_block(ws, "I/O LIST (from drawing PH-LCP-01 V1, sheets 10–13)", 6)
hdr = ["Address", "Tag", "Description", "Terminal", "Signal", "Sheet"]
for i, h in enumerate(hdr, 1):
    put(ws, 4, i, h, bold=True, fill=HDR, color="FFFFFF", h="center")
io = [
    ("I0.0", "PSH-101", "Scrubber high pressure", "XI:1", "DI 24V", "10"),
    ("I0.1", "LSH-102", "Scrubber high level", "XI:3", "DI 24V", "10"),
    ("I0.2", "LSHH-103", "Scrubber high-high level", "XI:5", "DI 24V", "10"),
    ("I0.3", "RR", "BR01 run feedback", "XI:7", "DI 24V", "10"),
    ("I0.4", "S1", "BR01 START push button", "XI:9", "DI 24V", "10"),
    ("I0.5", "Rflt", "BR01 fault (lockout)", "XI:11", "DI 24V", "10"),
    ("I0.6", "S3", "BR01 STOP push button", "XI:13", "DI 24V", "10"),
    ("I0.7", "—", "Spare", "XI:15", "DI 24V", "10"),
    ("I1.0", "S0", "E-STOP status (NC1)", "XI:17", "DI 24V", "10"),
    ("I1.1", "S2", "FAULT RESET push button", "XI:19", "DI 24V", "10"),
    ("I1.2", "LSL-TK01", "TK01 low level", "XI:21", "DI 24V", "10"),
    ("I1.3", "LSL-TK02", "TK02 low level", "XI:23", "DI 24V", "10"),
    ("I1.4", "RR2", "BR02 run feedback", "XI:25", "DI 24V", "10"),
    ("I1.5", "Rflt2", "BR02 fault (lockout)", "XI:27", "DI 24V", "10"),
    ("I12.0", "Rfuel2", "BR02 diesel mode", "XI2:1", "DI 24V", "11"),
    ("I12.1", "S4", "BR01 AUTO/MAN selector (closed = AUTO)", "XI2:3", "DI 24V", "11"),
    ("I12.2", "S5", "BR02 AUTO/MAN selector (closed = AUTO)", "XI2:5", "DI 24V", "11"),
    ("I12.3", "S6", "BR02 START push button", "XI2:7", "DI 24V", "11"),
    ("I12.4", "S7", "BR02 STOP push button", "XI2:9", "DI 24V", "11"),
    ("I12.5", "—", "Reserved: scrubber LSL (not installed)", "XI2:11", "DI 24V", "11"),
    ("I12.6–7", "—", "Spare", "XI2:13–16", "DI 24V", "11"),
    ("Q0.0", "R1", "BR01 heat request (TL) + H1 1st stage lamp", "XQ:1/2", "DQ 24V", "12"),
    ("Q0.1", "R2", "BR01 2nd stage (TR) + H2 2nd stage lamp", "XQ:3/4", "DQ 24V", "12"),
    ("Q0.2", "R3", "XV-401 Flare SV (NO)", "XQ:5/6", "DQ 24V", "12"),
    ("Q0.3", "R4", "XV-402 Drain SV (NO)", "XQ:7/8", "DQ 24V", "12"),
    ("Q0.4", "R5", "XV-403 Suction SV (NC)", "XQ:9/10", "DQ 24V", "12"),
    ("Q0.5", "R6", "BR02 start (L)", "X6:2", "DQ 24V", "12 / 8"),
    ("Q0.6", "R7", "BR02 high fire (T2)", "X6:4", "DQ 24V", "12 / 8"),
    ("Q0.7", "R8", "BR02 1st stage lamp H5", "—", "DQ 24V", "12"),
    ("Q1.0", "R9", "BR02 2nd stage lamp H6", "—", "DQ 24V", "12"),
    ("Q1.1", "R10", "Spare", "XQ:19/20", "DQ 24V", "12"),
    ("IW96", "TE-501", "TK01 inlet oil temperature", "X5:1/2", "TC type K", "13"),
    ("IW98", "TE-502", "TK01 outlet oil temperature", "X5:3/4", "TC type K", "13"),
    ("IW100", "TE-503", "TK02 inlet oil temperature", "X5:5/6", "TC type K", "13"),
    ("IW102", "TE-504", "TK02 outlet oil temperature", "X5:7/8", "TC type K", "13"),
    ("IW104–110", "—", "Spare thermocouple channels", "X5:9–16", "TC", "13"),
]
for k, row in enumerate(io, 5):
    for i, v in enumerate(row, 1):
        put(ws, k, i, v, h="center" if i in (1, 4, 5, 6) else "left")
for i, w in enumerate([11, 12, 42, 12, 10, 8], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A5"

# Drawing V2 changes are internal — they live in docs/ISSUES.md (F4–F7), not in the client file.

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    s.page_setup.horizontalCentered = True
    if s.title != "C&E Matrix":
        # one page wide, as many pages tall as needed
        s.page_setup.orientation = "landscape" if s.title == "Cover" else "portrait"
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0
        s.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("saved", OUT)
